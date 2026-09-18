import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_document_styles(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(0)

def add_major_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    return p

def add_p(doc, text, spacing=2.0, align=WD_PARAGRAPH_ALIGNMENT.LEFT, italic=False, bold=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.bold = bold
    return p

def add_single_image(doc, img_path, caption, width_in=5.0):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_img.paragraph_format.line_spacing = 1.0
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_cap.paragraph_format.line_spacing = 1.15
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True

def add_paired_images(doc, img1_path, img2_path, caption, width_in=2.3):
    if os.path.exists(img1_path) and os.path.exists(img2_path):
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in tbl.rows:
            for cell in row.cells:
                tcPr = cell._element.get_or_add_tcPr()
                tcBorders = parse_xml(r'<w:tcBorders %s><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>' % nsdecls('w'))
                tcPr.append(tcBorders)

        cell1 = tbl.cell(0, 0)
        p1 = cell1.paragraphs[0]
        p1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p1.paragraph_format.line_spacing = 1.0
        p1.add_run().add_picture(img1_path, width=Inches(width_in))

        cell2 = tbl.cell(0, 1)
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p2.paragraph_format.line_spacing = 1.0
        p2.add_run().add_picture(img2_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_cap.paragraph_format.line_spacing = 1.15
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True

def add_table_data(doc, title, headers, data):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run(title)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(11)
    run_t.font.bold = True

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = True
        shd = parse_xml(r'<w:shd %s w:fill="E8E8E8"/>' % nsdecls('w'))
        hdr_cells[i]._element.get_or_add_tcPr().append(shd)

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.line_spacing = 1.0
    p_spacer.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shd = parse_xml(r'<w:shd %s w:fill="F4F4F4"/>' % nsdecls('w'))
    cell._element.get_or_add_tcPr().append(shd)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_str)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 30, 30)

print("Thesis generator base initialized.")
