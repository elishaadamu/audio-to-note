import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Section 1: Preliminaries
sec1 = doc.sections[0]
sec1.different_first_page_header_footer = True  # Title page has no header/footer
pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman"/>')
sec1._sectPr.append(pgNumType1)

footer1 = sec1.footer
p1 = footer1.paragraphs[0]
p1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run1 = p1.add_run()
fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
run1._r.append(fld1)

doc.add_paragraph("Title Page")
doc.add_page_break()
doc.add_paragraph("Declaration (Page ii)")
doc.add_page_break()
doc.add_paragraph("Dedication (Page iii)")

# Section 2: Chapters
sec2 = doc.add_section()
sec2.footer.is_linked_to_previous = False
pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>')
sec2._sectPr.append(pgNumType2)

footer2 = sec2.footer
p2 = footer2.paragraphs[0]
p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run2 = p2.add_run()
fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
run2._r.append(fld2)

doc.add_paragraph("Chapter One (Page 1)")
doc.add_page_break()
doc.add_paragraph("Chapter One Continued (Page 2)")

doc.save("test_pages.docx")
print("Successfully generated test_pages.docx")
