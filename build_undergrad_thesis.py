import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_document_styles(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(0)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    return p

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_single_image(doc, img_path, caption, width_in=5.0):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_img.paragraph_format.line_spacing = 1.0
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_cap.paragraph_format.line_spacing = 1.15
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True

def add_paired_images(doc, img1_path, img2_path, caption, width_in=2.3):
    if os.path.exists(img1_path) and os.path.exists(img2_path):
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in tbl.rows:
            for cell in row.cells:
                tcPr = cell._element.get_or_add_tcPr()
                tcBorders = parse_xml(r'<w:tcBorders %s><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>' % nsdecls('w'))
                tcPr.append(tcBorders)

        cell1 = tbl.cell(0, 0)
        p1 = cell1.paragraphs[0]
        p1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p1.paragraph_format.line_spacing = 1.0
        p1.add_run().add_picture(img1_path, width=Inches(width_in))

        cell2 = tbl.cell(0, 1)
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p2.paragraph_format.line_spacing = 1.0
        p2.add_run().add_picture(img2_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_cap.paragraph_format.line_spacing = 1.15
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True

def add_table_data(doc, title, headers, data):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run(title)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(11)
    run_t.font.bold = True

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = True
        shd = parse_xml(r'<w:shd %s w:fill="E8E8E8"/>' % nsdecls('w'))
        hdr_cells[i]._element.get_or_add_tcPr().append(shd)

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.line_spacing = 1.0
    p_spacer.paragraph_format.space_after = Pt(6)

def generate_full_document():
    doc = Document()
    set_document_styles(doc)

    base_dir = r"c:\Users\Adamu\OneDrive\Documents\audio-to-note"
    student_img = r"C:\Users\Adamu\.gemini\antigravity-ide\brain\5cf4f373-f504-4e59-8ab1-ec4f0f91c062\nigerian_student_lecture_1788464296986.jpg"
    ch1_infographic = r"C:\Users\Adamu\.gemini\antigravity-ide\brain\5cf4f373-f504-4e59-8ab1-ec4f0f91c062\system_infographic_ch1_1788464307411.jpg"
    ch2_tech_infographic = r"C:\Users\Adamu\.gemini\antigravity-ide\brain\5cf4f373-f504-4e59-8ab1-ec4f0f91c062\tech_stack_infographic_ch2_1788464327506.jpg"
    react_img = os.path.join(base_dir, "assets", "images", "react-logo@3x.png")
    
    s1 = os.path.join(base_dir, "app-screenshots", "s1.jpg")
    s2 = os.path.join(base_dir, "app-screenshots", "s2.jpg")
    s3 = os.path.join(base_dir, "app-screenshots", "s3.jpg")
    s4 = os.path.join(base_dir, "app-screenshots", "s4.jpg")
    s5 = os.path.join(base_dir, "app-screenshots", "s5.jpg")
    s6 = os.path.join(base_dir, "app-screenshots", "s6.jpg")
    s7 = os.path.join(base_dir, "app-screenshots", "s7.jpg")
    s8 = os.path.join(base_dir, "app-screenshots", "s8.jpg")

    # =========================================================================
    # CHAPTER ONE: INTRODUCTION (Target: 3 to 4 pages with images)
    # =========================================================================
    add_title(doc, "CHAPTER ONE\nINTRODUCTION")

    add_heading_1(doc, "1.1 BACKGROUND OF THE STUDY")
    add_p(doc, 
        "In modern tertiary institutions, lectures remain the primary method of teaching and knowledge dissemination. "
        "Every day, university students attend various lectures where lecturers explain complex theories, solve mathematical "
        "problems, and share important academic concepts verbally. However, the traditional process of manual note-taking "
        "places a heavy burden on students. Trying to listen attentively, understand new concepts, and write down notes simultaneously "
        "is very difficult. Research in education shows that students often miss between 30% and 50% of the important details "
        "spoken during a fast-paced lecture."
    )
    add_p(doc, 
        "This challenge is especially noticeable in large Nigerian university lecture halls. In many Nigerian universities, "
        "classes can have between 200 and 800 students sitting together in large auditoriums or lecture theaters. Factors such as "
        "acoustic echoes, distance from the lecturer, spinning ceiling fans, and generator noise make it easy to miss crucial "
        "explanations. While some students attempt to record the audio using basic voice recording apps on their mobile phones, "
        "listening back to a two-hour audio recording before an exam is tedious, time-consuming, and difficult to search."
    )
    
    # Figure 1.1: Infographic of Audio-to-Note Concept
    add_single_image(doc, ch1_infographic, 
        "Figure 1.1: High-Level Conceptual Overview of the AI Audio-to-Note Converter System", 
        width_in=5.2
    )

    add_p(doc, 
        "Recent advancements in Artificial Intelligence (AI) and Machine Learning have created exciting opportunities to solve this "
        "problem. With modern Speech-to-Text (STT) technologies and Multimodal Large Language Models like Google's Gemini 2.5 Flash, "
        "it is now possible to automatically capture audio from a mobile phone, transcribe spoken words into text, and organize the "
        "text into clear, structured study notes. This project focuses on designing and implementing AudioNote AI, an intelligent "
        "mobile application that transforms spoken lecture audio into well-organized notes, verbatim transcripts, self-assessment "
        "quizzes, and multilingual translations."
    )

    # Figure 1.2: Nigerian student in lecture hall
    add_single_image(doc, student_img, 
        "Figure 1.2: Real-World Usage Scenario – A Nigerian University Student Utilizing the Mobile Audio-to-Note App During a Lecture", 
        width_in=4.6
    )

    add_heading_1(doc, "1.2 STATEMENT OF THE PROBLEM")
    add_p(doc, 
        "Manual note-taking in university lectures is inefficient and stressful for students. When students focus on writing down "
        "everything a lecturer says, they lose track of the core explanation and fail to internalize the topic being taught. Conversely, "
        "if they focus only on listening, they leave the lecture hall without detailed written reference materials for revision."
    )
    add_p(doc, 
        "Existing commercial transcription applications do not fully solve this issue for students. Most existing apps only provide a flat, "
        "unstructured wall of raw text without headings, bullet points, or summaries. Furthermore, they do not include student-centered "
        "learning features such as practice quizzes or translations into local Nigerian languages like Hausa, Yoruba, and Igbo. "
        "There is a clear need for an affordable, student-friendly mobile application that captures spoken audio and automatically produces "
        "clean study summaries, timestamped transcripts, and interactive study aids."
    )

    add_heading_1(doc, "1.3 AIM AND OBJECTIVES")
    add_p(doc, 
        "The aim of this project is to develop an AI-powered Audio-to-Note study companion that allows students to record lecture "
        "audio and automatically convert it into structured study notes, timestamped transcripts, self-assessment quizzes, and multilingual "
        "translations. To achieve this aim, the following specific objectives were formulated:"
    )
    add_p(doc, "1. To design and build an intuitive, responsive mobile user interface using React Native and Expo that allows students to easily record, pause, and review lecture audio.")
    add_p(doc, "2. To implement a backend server using Node.js, Express.js, and TypeScript to manage audio uploads, user accounts, and database storage securely.")
    add_p(doc, "3. To integrate Google's Gemini 2.5 Flash multimodal AI API to process recorded audio files, generate structured Markdown notes, and extract key points.")
    add_p(doc, "4. To implement an automated quiz generation feature that creates 10-question multiple-choice practice quizzes from the lecture content for active recall revision.")
    add_p(doc, "5. To provide a one-tap translation engine that converts study notes into 9 languages, including major Nigerian languages (Hausa, Igbo, and Yoruba).")
    add_p(doc, "6. To test and evaluate the application's transcription accuracy, ease of use, and response time in typical classroom conditions.")

    add_heading_1(doc, "1.4 SCOPE AND LIMITATIONS OF THE STUDY")
    add_p(doc, 
        "The scope of this project covers the software development of the AudioNote AI platform. It includes building the frontend mobile "
        "app using React Native, developing the REST API using Express.js and TypeScript, and managing persistent data with PostgreSQL and "
        "Prisma ORM. The artificial intelligence processing relies on Google's Gemini 2.5 Flash API for speech processing, note generation, "
        "and translations."
    )
    add_p(doc, 
        "The project has some limitations. First, because the audio transcription and AI processing are performed in the cloud, an active "
        "internet connection is required to process recorded lectures. Second, the clarity of the generated notes depends on the recording "
        "quality of the student's mobile phone; if a lecturer speaks very far from the microphone in a room with severe acoustic noise, some words "
        "may be misunderstood. Finally, the system relies on Google's API service availability and rate limits."
    )

    add_heading_1(doc, "1.5 SIGNIFICANCE OF THE STUDY")
    add_p(doc, 
        "This project is of significant benefit to students, educators, and the broader field of applied computer science. For students, it removes "
        "the pressure of multitasking during lectures, allowing them to engage actively in class discussions while trusting the app to capture the "
        "notes accurately. The automatic generation of interactive quizzes provides an effective self-assessment tool, helping students test their "
        "understanding before semester examinations."
    )
    add_p(doc, 
        "Additionally, offering note translations into indigenous languages like Hausa, Yoruba, and Igbo promotes inclusive education, helping "
        "students who understand concepts better in their native tongues. From a computer science perspective, the project demonstrates how modern "
        "web and mobile technologies (React Native, Node.js, and PostgreSQL) can be integrated with state-of-the-art multimodal AI services to "
        "solve practical educational problems in developing regions."
    )

    add_heading_1(doc, "1.6 OPERATIONAL DEFINITION OF TERMS")
    add_p(doc, "• Speech-to-Text (STT): The technology that converts spoken words into computer-readable digital text.")
    add_p(doc, "• Multimodal AI: An artificial intelligence system capable of understanding and processing different types of input data, such as audio, text, and images simultaneously.")
    add_p(doc, "• React Native: A popular open-source JavaScript framework developed by Meta for building natively rendering mobile applications for Android and iOS from a single codebase.")
    add_p(doc, "• Expo: A set of tools, libraries, and services built around React Native that simplifies mobile app development, testing, and deployment.")
    add_p(doc, "• Express.js: A lightweight and flexible Node.js web application framework used to build backend RESTful APIs.")
    add_p(doc, "• Prisma ORM: An open-source database toolkit that provides type-safe database queries and automated schema migrations for PostgreSQL.")
    add_p(doc, "• Active Recall: A study technique where learners actively test their memory through quizzes and questions rather than passively re-reading notes.")

    # Page break between Chapter 1 and Chapter 2
    doc.add_page_break()

    # =========================================================================
    # CHAPTER TWO: LITERATURE REVIEW (Target: ~21 pages with images included)
    # =========================================================================
    add_title(doc, "CHAPTER TWO\nLITERATURE REVIEW")

    add_heading_1(doc, "2.0 OVERVIEW OF AI-POWERED AUDIO-TO-NOTE SYSTEMS")
    add_p(doc, 
        "Converting spoken dialogue into written text has always been a major goal of computer science. In academic environments, lectures "
        "contain valuable knowledge that students need to preserve and study. Early attempts to capture lectures relied entirely on manual "
        "note-taking or physical audio cassette recordings. However, raw audio files are difficult to review because users cannot search "
        "through them quickly. Students must spend hours listening to find specific definitions or explanations."
    )
    add_p(doc, 
        "With the rapid growth of Artificial Intelligence (AI) and Natural Language Processing (NLP), modern software systems can do much "
        "more than simple audio playback. Today's AI systems can listen to human speech, convert the audio into text, understand the context of "
        "what was said, and automatically organize the information into clear summaries. This chapter provides a detailed review of the literature "
        "surrounding Speech-to-Text (STT) technologies, Natural Language Processing for text summarization, multimodal AI models, mobile "
        "application architectures, and existing audio note-taking systems."
    )

    # Insert Chapter 2 Architecture Infographic
    add_single_image(doc, ch2_tech_infographic, 
        "Figure 2.1: Full-Stack Architecture Diagram Illustrating the Frontend, Backend, AI Pipeline, and Database Layers", 
        width_in=5.4
    )

    add_heading_1(doc, "2.1 SPEECH-TO-TEXT (STT) TECHNOLOGIES AND PROCESSING METHODS")
    add_p(doc, 
        "Automatic Speech Recognition (ASR), commonly known as Speech-to-Text (STT), is the technology that enables a computer program to "
        "recognize human speech and convert it into written words. The speech recognition process involves receiving an acoustic wave from a "
        "microphone, digitizing the signal, analyzing sound frequencies, and decoding the sound patterns into written language."
    )

    add_heading_2(doc, "2.1.1 Types of Audio Inputs")
    add_p(doc, 
        "Speech-to-text systems encounter different types of audio inputs depending on the recording environment and hardware. In a laboratory "
        "setting, audio is often recorded with close-talking microphones in quiet rooms, producing clean sound with high signal quality. "
        "However, in real-world educational settings, audio inputs are far more challenging."
    )
    add_p(doc, 
        "Audio inputs can be broadly categorized into two types: real-time streaming audio and pre-recorded audio files. Real-time streaming "
        "involves processing small audio chunks (typically 250 milliseconds to a few seconds) as the speaker is talking. This requires low-latency "
        "processing so that the user sees words appear almost immediately. Pre-recorded audio files, on the other hand, allow the system to receive "
        "the entire recording at once. This enables the AI model to analyze the full context of the speech from start to finish, which generally "
        "leads to higher transcription accuracy."
    )

    add_heading_2(doc, "2.1.2 Speech Recognition Methods")
    add_p(doc, 
        "Over the decades, speech recognition systems have evolved through several technological generations. Understanding these methods "
        "explains why modern deep learning systems are significantly more capable than older approaches."
    )

    add_heading_3(doc, "2.1.2.1 Acoustic Modeling")
    add_p(doc, 
        "The acoustic model is the part of an ASR system that connects audio signals to the basic units of speech, known as phonemes. In English, "
        "there are approximately 44 distinct phonemes that make up spoken words. In early systems, researchers used Gaussian Mixture Models (GMM) "
        "combined with Hidden Markov Models (HMM). The GMM estimated the probability of a sound belonging to a specific phoneme, while the HMM "
        "tracked the sequence of sounds over time."
    )
    add_p(doc, 
        "While statistical HMMs were an important milestone, they struggled with background noise and speaker variations. Modern speech recognition "
        "replaces GMMs with Deep Neural Networks (DNNs) and Convolutional Neural Networks (CNNs). Deep learning models are capable of learning "
        "complex acoustic patterns directly from thousands of hours of training audio, making them far more resilient to different voices, pitches, "
        "and accents."
    )

    add_heading_3(doc, "2.1.2.2 Language Modeling")
    add_p(doc, 
        "While the acoustic model identifies what sounds were spoken, the language model determines which words make sense together grammatically "
        "and contextually. Human speech contains many homophones—words that sound identical but have different spellings and meanings, such as "
        "'their', 'there', and 'they're', or 'hear' and 'here'. Without a language model, the computer cannot know which word is correct."
    )
    add_p(doc, 
        "Early language models relied on statistical n-grams, which calculated the probability of a word appearing based on the previous two or "
        "three words. Today, language modeling is powered by Transformer neural networks, such as those found in modern Large Language Models. "
        "Transformers use self-attention mechanisms to look at the entire sentence or paragraph, allowing the system to easily select the correct "
        "word based on the overall meaning of the lecture."
    )

    # Figure 2.2: Audio Capture Initiation & Dynamic Waveform
    add_paired_images(doc, s2, s3, 
        "Figure 2.2: Audio Recording Interfaces in AudioNote AI – (Left: Standby Recording Screen [s2]; Right: Active Waveform Metering [s3])",
        width_in=2.3
    )

    add_heading_2(doc, "2.1.3 Audio Processing and Feature Analysis Techniques")
    add_p(doc, 
        "Before an audio signal can be analyzed by a machine learning model, it must be cleaned and converted into a numerical format that computers "
        "can process efficiently."
    )

    add_heading_3(doc, "2.1.3.1 Pre-Processing and Noise Reduction")
    add_p(doc, 
        "Raw audio captured from a phone microphone often contains unwanted noise such as electrical hum, wind, or distant talking. Pre-processing "
        "applies digital filters to clean the audio signal. A common step is pre-emphasis filtering, which boosts high-frequency sounds that naturally "
        "drop off when humans speak. Another important technique is noise suppression, which detects steady background noise (such as an air conditioner "
        "or electric generator) and subtracts that noise from the audio signal, leaving a clearer voice track."
    )

    add_heading_3(doc, "2.1.3.2 Feature Extraction (Spectrograms and MFCCs)")
    add_p(doc, 
        "Computers cannot directly process raw sound waves efficiently because sound waves are simply long arrays of amplitude values over time. "
        "To make sense of sound, systems extract Mel-Frequency Cepstral Coefficients (MFCCs) or compute Log-Mel Spectrograms. A spectrogram is a "
        "visual representation of sound frequencies over time. The 'Mel scale' is used because it models how human ears perceive pitch—humans are "
        "much better at distinguishing small differences in low frequencies than in high frequencies. By converting audio into spectrograms, "
        "computer vision and deep learning techniques can treat the audio almost like an image to recognize words accurately."
    )

    add_heading_3(doc, "2.1.3.3 Speaker Diarization")
    add_p(doc, 
        "Speaker diarization is the process of identifying 'who spoke when' in an audio recording. In a classroom, a lecture is not always a monologue; "
        "students frequently ask questions or participate in class discussions. Diarization algorithms analyze voice characteristics (such as pitch "
        "and vocal tone) and assign labels like 'Speaker 1' (the lecturer) and 'Speaker 2' (a student asking a question). This helps organize the "
        "resulting transcript into a natural dialogue format."
    )

    # Insert Table 2.1: Comparison of ASR Approaches
    add_table_data(doc, 
        "Table 2.1: Comparison of Speech Recognition Approaches in Computer Science",
        ["Approach", "Technology", "Accuracy", "Handling of Accents", "Computational Needs"],
        [
            ["Statistical ASR", "Hidden Markov Models (HMM) + GMM", "Low to Moderate (70-80%)", "Poor; struggles with variations", "Low CPU usage; runs on old PCs"],
            ["Hybrid Deep Learning", "Deep Neural Networks + HMM", "Moderate to High (85-90%)", "Fair; requires accent tuning", "Moderate; requires decent CPU/GPU"],
            ["End-to-End ASR", "Connectionist Temporal Classification (CTC)", "High (90-95%)", "Good across standard dialects", "High GPU requirements for training"],
            ["Multimodal Foundation AI", "Transformer-based (e.g. Gemini 2.5 Flash)", "Very High (95-98%+)", "Excellent across global & regional accents", "Offloaded to cloud AI clusters"]
        ]
    )

    add_heading_1(doc, "2.2 NATURAL LANGUAGE PROCESSING (NLP) FOR SUMMARIZATION")
    add_p(doc, 
        "Getting an accurate transcript is only the first step. A word-for-word transcript of a one-hour lecture can easily be over 8,000 words long. "
        "Reading through 8,000 words of spoken dialogue is overwhelming because human speech includes filler words (like 'um', 'you know', 'actually'), "
        "repetitive explanations, and side remarks. Natural Language Processing (NLP) summarization techniques are needed to distill this raw text "
        "into clear, concise study notes."
    )

    add_heading_2(doc, "2.2.1 Extractive Summarization")
    add_p(doc, 
        "Extractive summarization works by identifying the most important sentences already present in the text and stitching them together into a "
        "summary. Algorithms such as TF-IDF (Term Frequency-Inverse Document Frequency) and TextRank score each sentence based on the frequency and "
        "importance of its words. While extractive summarization is fast and guaranteed not to make up facts, it often creates choppy, disjointed summaries "
        "because spoken sentences do not always flow logically when pulled out of context."
    )

    add_heading_2(doc, "2.2.2 Abstractive Summarization")
    add_p(doc, 
        "Abstractive summarization is much closer to how a human student writes notes. Instead of copying exact sentences, an abstractive AI model "
        "reads the transcript, understands the main concepts, and writes completely new, well-structured sentences. This allows the model to "
        "eliminate verbal clutter, combine related points, create clean bulleted lists, and format the output with markdown headings. Modern "
        "Large Language Models (LLMs) excel at abstractive summarization because they have been trained on vast amounts of literature and know how "
        "to synthesize academic concepts clearly."
    )

    add_heading_2(doc, "2.2.3 Keyword and Topic Extraction")
    add_p(doc, 
        "To make study notes easy to organize and search, NLP systems perform keyword and topic extraction. Using techniques like Named Entity "
        "Recognition (NER), the system automatically detects technical terms, formulas, dates, and author names. In AudioNote AI, this allows the "
        "app to generate automatic titles (such as 'Full-Stack Development: Front-End, Back-End, and Hosting') and tag notes by subject categories "
        "without requiring the student to type anything manually."
    )

    add_heading_2(doc, "2.2.4 Challenges in Automated Summarization")
    add_p(doc, 
        "While abstractive summarization is powerful, it comes with challenges. The main challenge is 'hallucination', where an AI model generates "
        "information that sounds convincing but was never mentioned in the original audio. In an academic setting, hallucinating an incorrect formula "
        "or wrong historical date can misinform students. Developers address this by carefully engineering prompts with strict rules that tell the AI "
        "to base its summary solely on the provided audio transcript."
    )

    # Figure 2.3: Decision Screen and Processed Note View
    add_paired_images(doc, s4, s6, 
        "Figure 2.3: Ingestion and Note Generation – (Left: Saved Recording Decision Screen [s4]; Right: Processed Note with Gemini 2.5 Flash Badge [s6])",
        width_in=2.3
    )

    add_heading_1(doc, "2.3 REAL-WORLD CHALLENGES IN AUDIO CONVERSION")
    add_p(doc, 
        "Building a reliable speech-to-note application for real university environments requires overcoming several practical challenges. "
        "These challenges are particularly pronounced in developing educational contexts such as Nigerian higher institutions."
    )

    add_heading_2(doc, "2.3.1 Background Noise and Acoustic Reflections")
    add_p(doc, 
        "In many universities, lecture rooms are large concrete halls without acoustic wall padding. When a lecturer speaks, the sound reflects "
        "off walls, floors, and metal desks, creating reverberation (echo). Additionally, environmental sounds such as rain hitting corrugated iron "
        "roofs, diesel generator noise during power outages, noisy ceiling fans, and chatter among students can drown out the lecturer's voice. "
        "If the audio recorded on the student's phone is muffled, speech recognition systems face higher error rates."
    )

    add_heading_2(doc, "2.3.2 Accents and Regional Dialectal Variations")
    add_p(doc, 
        "English is the official language of instruction in Nigeria, but it is spoken with a wide variety of regional accents influenced by the "
        "speaker's mother tongue (such as Yoruba, Hausa, or Igbo). Many commercial speech recognition tools developed in Western countries were "
        "trained predominantly on North American or British accents, causing them to struggle with Nigerian English pronunciations. Furthermore, "
        "lecturers frequently use Nigerian idioms or switch briefly into pidgin or local languages to explain a concept. A robust educational app "
        "needs an AI model with broad global language exposure to understand these speech variations."
    )

    add_heading_2(doc, "2.3.3 Domain-Specific Jargon and Technical Terms")
    add_p(doc, 
        "Computer science, engineering, and medical lectures are filled with specialized terminology. For example, a computer science lecturer "
        "might mention 'PostgreSQL', 'Kubernetes', 'asynchronous functions', or 'polymorphism'. A general speech-to-text model that has not been "
        "conditioned for technical topics might transcribe 'PostgreSQL' as 'post grey sequel' or 'API' as 'a pie'. Ensuring that technical terms "
        "are transcribed accurately requires intelligent language models that recognize the academic context."
    )

    add_heading_2(doc, "2.3.4 Mobile Hardware and Data Network Constraints")
    add_p(doc, 
        "Undergraduate students use a wide variety of mobile smartphones, ranging from budget Android phones with modest processors to modern devices. "
        "An application cannot expect to run a massive AI model locally on the phone because it would quickly drain the battery and overheat the device. "
        "At the same time, mobile data in Nigeria can be costly and campus Wi-Fi may be intermittent. Therefore, the mobile application must compress "
        "audio efficiently before sending it to the cloud server, minimizing data usage while preserving audio clarity."
    )

    # Insert Table 2.2: Comparison of Text Summarization Techniques
    add_table_data(doc, 
        "Table 2.2: Comparison of Text Summarization Techniques for Lecture Notes",
        ["Summarization Type", "Technique Used", "Readability", "Risk of Hallucination", "Suitability for Study Notes"],
        [
            ["Sentence Extraction", "TF-IDF / Frequency Scoring", "Low; disjointed sentences", "None (copies verbatim)", "Poor; keeps verbal filler and pauses"],
            ["Graph-Based Extraction", "TextRank Algorithm", "Moderate; selects key sentences", "None (copies verbatim)", "Fair; lacks natural conversational flow"],
            ["Early Seq2Seq Neural", "LSTM / RNN Encoder-Decoder", "Moderate; some grammatical errors", "Low to Moderate", "Fair; forgets context in long lectures"],
            ["Modern Abstractive AI", "Transformer LLMs (Gemini 2.5 Flash)", "Very High; structured markdown", "Low with proper prompt grounding", "Excellent; produces clear academic summaries"]
        ]
    )

    add_heading_1(doc, "2.4 MACHINE LEARNING APPROACHES IN SPEECH RECOGNITION")
    add_p(doc, 
        "Machine learning is the branch of computer science that allows systems to learn from data rather than following strictly hardcoded rules. "
        "In speech processing, machine learning models are trained on thousands of hours of audio recordings paired with text transcripts to learn "
        "the mapping from sound vibrations to written words."
    )

    add_heading_2(doc, "2.4.1 Supervised, Unsupervised, and Self-Supervised Learning")
    add_p(doc, 
        "Traditionally, speech recognition relied on supervised learning, where every second of training audio had to be manually transcribed "
        "by humans. This was extremely expensive and made it difficult to build speech models for languages with fewer resources. Recently, "
        "the AI community adopted self-supervised learning (SSL) models, such as wav2vec 2.0. In self-supervised learning, the AI model listens to "
        "massive amounts of unlabeled audio and learns how speech sounds work on its own by trying to predict masked or hidden parts of the audio. "
        "Once this general understanding is established, the model only needs a small amount of labeled text to become an accurate transcriber."
    )

    add_heading_2(doc, "2.4.2 The Transformer Architecture in Speech and NLP")
    add_p(doc, 
        "The introduction of the Transformer architecture by Vaswani et al. in 2017 fundamentally changed natural language processing and speech "
        "recognition. Before Transformers, systems processed words one by one using recurrent loops (RNNs), which were slow and frequently forgot "
        "information from earlier parts of long sentences. The Transformer introduced the 'self-attention' mechanism, which allows the model to "
        "examine all parts of an audio clip or text document at the same time. This capability allows models to understand the broad context of an "
        "entire lecture and generate accurate notes without losing the central topic."
    )

    # Figure 2.4: Synchronized Playback and Multilingual Translation Modal
    add_paired_images(doc, s7, s8, 
        "Figure 2.4: Study and Accessibility Features – (Left: Synchronized Audio Playback [s7]; Right: Multilingual Translation Dialog [s8])",
        width_in=2.3
    )

    add_heading_1(doc, "2.5 ADVANCEMENTS IN MULTIMODAL AI AND PEDAGOGICAL TOOLS")
    add_p(doc, 
        "The latest breakthrough in artificial intelligence is the shift towards native Multimodal Models. Earlier AI tools operated in separate "
        "stages: one model converted audio to text, another model summarized the text, and a third model translated it. If the first model made a "
        "mistake, the mistake propagated and grew worse through each subsequent stage. Native multimodal models eliminate this problem by processing "
        "audio directly within a single neural network."
    )

    add_heading_2(doc, "2.5.1 Google Gemini 2.5 Flash Multimodal Pipeline")
    add_p(doc, 
        "Google's Gemini 2.5 Flash is an advanced multimodal model designed for high-speed, cost-effective reasoning. Unlike older text-only models, "
        "Gemini 2.5 Flash can ingest audio files directly. It listens to the audio, understands the speaker's tone, emphasis, and context, and generates "
        "structured text in one step. In AudioNote AI, the backend uses Google's AI File Manager API to stream the lecture audio directly to Gemini. "
        "The model returns an organized JSON payload containing a clear lecture title, categorized topic tags, a structured markdown summary, "
        "a timestamped verbatim transcript, and a practice quiz."
    )

    add_heading_2(doc, "2.5.2 Automated Quiz Generation for Active Recall")
    add_p(doc, 
        "Educational research consistently shows that simply reading through notes is one of the least effective ways to study. Cognitive scientists "
        "recommend 'active recall'—testing oneself on the material to strengthen memory retention. By leveraging Gemini's generative capabilities, "
        "AudioNote AI automatically turns lecture content into an interactive 10-question multiple-choice quiz. Each question includes four answer "
        "options and explanations for why an answer is correct. In the mobile app, students can answer questions and receive instant scoring and "
        "celebratory feedback, turning study sessions into an engaging learning experience."
    )

    add_heading_2(doc, "2.5.3 Multilingual Translation for Inclusive Learning")
    add_p(doc, 
        "Language accessibility is vital for students who are non-native English speakers. In Nigeria, while university education is in English, "
        "many students find it easier to understand difficult technical concepts when explained in their mother tongue. AudioNote AI integrates "
        "a neural translation engine supporting 9 languages: English, French, Spanish, German, Chinese, Arabic, and three prominent Nigerian languages: "
        "Hausa, Yoruba, and Igbo. With one tap, students can translate the summary and transcript into their chosen language while preserving "
        "all headings, bullet points, and timestamp markers."
    )

    # Insert Table 2.3: Multimodal Models Comparison
    add_table_data(doc, 
        "Table 2.3: State-of-the-Art Multimodal Foundation Models for Educational Applications",
        ["Model", "Provider", "Direct Audio Ingestion", "Speed / Latency", "Support for African Languages"],
        [
            ["Whisper Large-v3", "OpenAI", "Yes (Speech only)", "Moderate (generates text only)", "Moderate; limited for regional dialects"],
            ["GPT-4o Audio", "OpenAI", "Yes (Multimodal)", "Fast (high cloud API cost)", "Good global coverage"],
            ["Claude 3.5 Sonnet", "Anthropic", "No (requires text transcript first)", "Fast (text only)", "Good global coverage"],
            ["Gemini 2.5 Flash", "Google", "Yes (Direct Multimodal audio & text)", "Very Fast (low cost & high speed)", "Strong multilingual support including Nigerian languages"]
        ]
    )

    add_heading_1(doc, "2.6 MOBILE-FIRST ARCHITECTURE AND TECHNOLOGIES USED")
    add_p(doc, 
        "To deliver a smooth, practical experience to students, the software architecture of AudioNote AI was carefully designed using modern "
        "open-source JavaScript and TypeScript libraries. The system consists of a cross-platform mobile frontend, a RESTful backend API, and a "
        "relational database."
    )

    # Figure 2.5: React Native Logo
    add_single_image(doc, react_img, 
        "Figure 2.5: React Native Framework Architectural Ecosystem Utilized for the Mobile Client", 
        width_in=2.0
    )

    add_heading_2(doc, "2.6.1 React Native and Expo Framework")
    add_p(doc, 
        "React Native is a cross-platform mobile framework developed by Meta that allows developers to write code once in JavaScript or TypeScript "
        "and deploy it natively to both Android and iOS devices. Unlike hybrid web-view wrappers, React Native renders true native UI components, "
        "ensuring 60 frames-per-second performance and a smooth user experience. Expo (SDK 54) provides a comprehensive set of tools that simplify "
        "hardware access, app building, and testing without requiring complex native platform setup."
    )

    add_heading_2(doc, "2.6.2 Audio Recording and Waveform Libraries")
    add_p(doc, 
        "Capturing audio reliably on a smartphone requires specialized libraries. In this project, `expo-av` and `expo-audio` manage the microphone "
        "hardware sessions. To give the student visual feedback that their recording is working, the app polls the microphone's decibel levels "
        "every 250 milliseconds and animates a dynamic audio waveform using `react-native-reanimated`. The recorded audio is temporarily stored in "
        "`FileSystem.documentDirectory` on the phone before being uploaded to the backend server."
    )

    add_heading_2(doc, "2.6.3 Backend Architecture: Node.js, Express.js, and Prisma ORM")
    add_p(doc, 
        "The backend service is built using Node.js and Express.js in TypeScript. Express.js provides a clean routing mechanism for handling "
        "HTTP requests such as user registration, audio uploads, and note retrieval. Audio file uploads are handled using `multer`, which streams "
        "incoming audio directly to disk storage to avoid consuming excessive server RAM. Database operations are handled using Prisma ORM v6 "
        "connected to a PostgreSQL relational database. Prisma provides type-safe database queries, automated migrations, and connection pooling, "
        "ensuring that user records and generated notes are stored securely and retrieved quickly."
    )

    # Figure 2.6: PIN Unlock and User Settings
    add_paired_images(doc, s1, s5, 
        "Figure 2.6: Security and User Settings – (Left: 4-Digit Security PIN Gateway [s1]; Right: User Settings and Speech Model Configuration [s5])",
        width_in=2.3
    )

    add_heading_1(doc, "2.7 RELATED WORKS")
    add_p(doc, 
        "Several commercial and academic systems have explored audio transcription and note-taking. Examining these existing works helps "
        "highlight the unique contributions of AudioNote AI."
    )
    add_p(doc, 
        "Otter.ai is one of the most well-known automated meeting transcription platforms. It provides live transcription, speaker identification, "
        "and integration with video platforms like Zoom. However, Otter.ai is designed primarily for corporate business meetings. Its summaries "
        "focus on meeting action items rather than academic lecture concepts. In addition, its subscription costs are prohibitive for students in "
        "developing countries, and it does not support translations into local African languages or automated study quizzes."
    )
    add_p(doc, 
        "Fireflies.ai is another conversational intelligence tool focused on recording and transcribing corporate calls. It connects with business "
        "tools like Slack and Salesforce, but it lacks a dedicated mobile study interface with flashcards, active recall quizzes, or student-oriented "
        "note formatting. Other apps such as Coconote and Wave AI have recently emerged for students, but many of them rely on multi-step pipelines "
        "that introduce transcription errors and require expensive recurring monthly subscriptions."
    )
    add_p(doc, 
        "AudioNote AI addresses these gaps by combining direct multimodal AI processing (Google Gemini 2.5 Flash) with an undergraduate-focused "
        "mobile design. It provides structured lecture notes, timestamped transcripts, self-assessment quizzes, and translations into major Nigerian "
        "languages, offering a comprehensive, inclusive academic study companion."
    )

    add_heading_1(doc, "2.8 SUMMARY OF LITERATURE REVIEW")
    add_p(doc, 
        "In conclusion, the literature review shows that manual note-taking creates significant cognitive fatigue for students, causing them "
        "to miss important lecture material. While early speech recognition tools were limited by noise sensitivity and inflexible statistical models, "
        "modern deep learning and multimodal foundation models like Gemini 2.5 Flash have made high-accuracy speech-to-note conversion a reality. "
        "By leveraging React Native on the mobile client, Express.js on the backend, and Google's multimodal AI in the cloud, this project provides "
        "a practical, student-centered solution designed to overcome the acoustic and linguistic challenges found in university classrooms."
    )

    output_filename = "Audio_to_Note_Final_Thesis_Undergrad.docx"
    output_path = os.path.join(base_dir, output_filename)
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    generate_full_document()
