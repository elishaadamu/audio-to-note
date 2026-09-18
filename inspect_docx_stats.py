import docx

doc = docx.Document(r"c:\Users\Adamu\OneDrive\Documents\audio-to-note\Audio_to_Note_Converter_Final_Thesis.docx")

total_words = 0
ch1_words = 0
ch2_words = 0
is_ch2 = False

for p in doc.paragraphs:
    text = p.text.strip()
    if "CHAPTER TWO" in text:
        is_ch2 = True
    w = len(text.split())
    total_words += w
    if is_ch2:
        ch2_words += w
    else:
        ch1_words += w

table_words = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            table_words += len(cell.text.split())

print(f"Chapter 1 Words: {ch1_words}")
print(f"Chapter 2 Words: {ch2_words}")
print(f"Table Words: {table_words}")
print(f"Total Words: {total_words + table_words}")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
